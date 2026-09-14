import asyncio
import inspect
import os
from pathlib import Path

import pytest

from server.catalog import configure_import_paths, load_operation, operation_by_name


ROOT = Path(os.environ["ARCHIPELAGO_ROOT"]) if os.environ.get("ARCHIPELAGO_ROOT") else None
pytestmark = pytest.mark.skipif(ROOT is None or not ROOT.exists(), reason="staged Archipelago source is unavailable")


def invoke(operation, payload):
    assert ROOT is not None
    configure_import_paths(ROOT, operation.domain)
    function, _ = load_operation(operation, ROOT)
    annotation = inspect.get_annotations(function, eval_str=True)
    parameter = next(iter(inspect.signature(function).parameters.values()), None)
    argument = annotation[parameter.name].model_validate(payload) if parameter and hasattr(annotation[parameter.name], "model_validate") else payload
    result = function(argument) if parameter and hasattr(annotation.get(parameter.name), "model_validate") else function(**payload)
    return asyncio.run(result) if inspect.isawaitable(result) else result


def test_document_create_and_reopen(tmp_path):
    os.environ["APP_FS_ROOT"] = str(tmp_path)
    invoke(operation_by_name("documents", "create_document"), {
        "directory": "/",
        "file_name": "smoke.docx",
        "content": [{"type": "paragraph", "text": "Office adapter smoke test"}],
    })
    assert (tmp_path / "smoke.docx").exists()
    content = invoke(operation_by_name("documents", "read_document_content"), {"file_path": "/smoke.docx"})
    assert "Office adapter smoke test" in str(content)


def test_spreadsheet_create_and_reopen(tmp_path):
    os.environ["APP_FS_ROOT"] = str(tmp_path)
    invoke(operation_by_name("spreadsheets", "create_spreadsheet"), {
        "directory": "/",
        "file_name": "smoke.xlsx",
        "sheets": [{"name": "Data", "headers": ["Value"], "rows": [[1]]}],
    })
    invoke(operation_by_name("spreadsheets", "add_content_text"), {
        "file_path": "/smoke.xlsx",
        "tab_index": 0,
        "cell": "A3",
        "value": "=SUM(A2:A2)",
    })
    from openpyxl import load_workbook

    workbook = load_workbook(tmp_path / "smoke.xlsx", data_only=False)
    assert workbook["Data"]["A2"].value == 1
    assert workbook["Data"]["A3"].value == "=SUM(A2:A2)"
    workbook.close()


def test_presentation_create_and_reopen(tmp_path):
    os.environ["APP_FS_ROOT"] = str(tmp_path)
    invoke(operation_by_name("presentations", "create_deck"), {
        "directory": "/",
        "file_name": "smoke.pptx",
        "slides": [{"layout": "title", "title": "Office adapter smoke test"}],
    })
    from pptx import Presentation

    presentation = Presentation(tmp_path / "smoke.pptx")
    assert len(presentation.slides) == 1
